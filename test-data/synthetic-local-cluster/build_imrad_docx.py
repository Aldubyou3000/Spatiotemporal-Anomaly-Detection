"""Reorganize recorded controlled evaluation evidence into an editable IMRaD document."""
import json
from pathlib import Path

from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parent
OUT = ROOT / "evidence/extended"


def main():
    metrics = json.loads((ROOT / "evidence/metrics.json").read_text())
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Inches(8.5), Inches(11)
    section.top_margin = section.bottom_margin = section.left_margin = section.right_margin = Inches(1)
    for name in ("Normal", "Title", "Subtitle", "Heading 1", "Heading 2", "Heading 3", "Caption"):
        style = doc.styles[name]
        style.font.name = "Times New Roman"
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.font.size = Pt(12)
        style.paragraph_format.space_after = Pt(6)
        if style.element.rPr is not None:
            for color in style.element.rPr.findall(qn("w:color")):
                color.attrib.pop(qn("w:themeColor"), None)
    doc.styles["Normal"].paragraph_format.line_spacing = 1.5
    doc.styles["Title"].font.size = Pt(16)
    doc.styles["Title"].font.bold = True
    doc.styles["Title"].paragraph_format.line_spacing = 1.15
    for name in ("Heading 1", "Heading 2", "Heading 3"):
        doc.styles[name].font.bold = True
        doc.styles[name].paragraph_format.space_before = Pt(12)
        doc.styles[name].paragraph_format.keep_with_next = True
    doc.styles["Caption"].font.size = Pt(10)
    doc.styles["Caption"].font.italic = True
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)
    doc.core_properties.title = "Controlled Testing of Rainfall Anomaly Detection"
    doc.core_properties.subject = "IMRaD report of synthetic controlled testing"
    doc.core_properties.author = ""
    doc.core_properties.keywords = "rainfall, anomaly detection, controlled evaluation, IMRaD"

    def p(text):
        paragraph = doc.add_paragraph(text)
        paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        return paragraph

    def heading(text, level=2):
        doc.add_heading(text, level=level)

    def caption(text):
        paragraph = doc.add_paragraph(text, style="Caption")
        paragraph.paragraph_format.keep_with_next = True

    def table(rows, widths):
        result = doc.add_table(rows=0, cols=len(rows[0]))
        result.alignment = WD_TABLE_ALIGNMENT.CENTER
        result.autofit = False
        for column, width in zip(result.columns, widths):
            column.width = Inches(width)
        for index, values in enumerate(rows):
            row = result.add_row()
            no_split = OxmlElement("w:cantSplit")
            row._tr.get_or_add_trPr().append(no_split)
            if index == 0:
                repeat = OxmlElement("w:tblHeader")
                row._tr.get_or_add_trPr().append(repeat)
            for number, (cell, value, width) in enumerate(zip(row.cells, values, widths)):
                cell.width = Inches(width)
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                properties = cell._tc.get_or_add_tcPr()
                margins = OxmlElement("w:tcMar")
                for side in ("top", "left", "bottom", "right"):
                    margin = OxmlElement(f"w:{side}")
                    margin.set(qn("w:w"), "90")
                    margin.set(qn("w:type"), "dxa")
                    margins.append(margin)
                properties.append(margins)
                borders = OxmlElement("w:tcBorders")
                for side in ("top", "left", "bottom", "right"):
                    border = OxmlElement(f"w:{side}")
                    border.set(qn("w:val"), "single")
                    border.set(qn("w:sz"), "4")
                    border.set(qn("w:color"), "D9D9D9")
                    borders.append(border)
                properties.append(borders)
                if index == 0:
                    fill = OxmlElement("w:shd")
                    fill.set(qn("w:fill"), "DCE6F1")
                    properties.append(fill)
                paragraph = cell.paragraphs[0]
                paragraph.paragraph_format.line_spacing = 1
                paragraph.paragraph_format.space_after = Pt(0)
                paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT if number == 0 else WD_ALIGN_PARAGRAPH.CENTER
                run = paragraph.add_run(str(value))
                run.font.size = Pt(10)
                run.bold = index == 0
        doc.add_paragraph().paragraph_format.space_after = Pt(2)

    title = doc.add_paragraph("Controlled Testing of Rainfall Anomaly Detection", "Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    heading("Introduction", 1)
    p("This test checked whether the system could identify unusual rainfall readings in a controlled dataset. "
      "The dataset used nearby fictional stations with readings on the same dates so that the system could compare them. "
      "The purpose was to measure correct detections, missed anomalies and false alarms.")
    heading("Methods", 1)
    p("Synthetic rainfall data were prepared for eight fictional stations. All stations covered September 1 to 30, 2023, "
      "with 24 hourly readings per day. The selected neighboring stations were approximately 0.43 to 0.62 kilometres apart. "
      "The system converted the hourly readings into 240 daily observations. Each observation represented one station on one day.")
    p("The dataset contained 16 deliberately introduced measurement errors: eight unusually high readings and eight unusually low readings. "
      "The remaining 224 observations were labelled normal. Normal cases included dry days, rainy days and valid differences between stations.")
    p("An answer key identified the intended normal and anomalous observations before the system was run. "
      "The system's high or low anomaly flags were compared with that key using station name and date. "
      "All 240 daily observations were assessed. The existing detection rules were used without changing their settings.")
    p("The comparison produced a confusion matrix. Precision, recall, F1 score and accuracy were then calculated from its counts. "
      "These measures describe the system's combined high and low flagging results.")
    heading("Results", 1)
    p("The system detected 12 of the 16 introduced anomalies and missed four. It also flagged seven normal observations as anomalies "
      "and correctly left 217 normal observations unflagged.")
    caption("Table 1 Confusion matrix for the controlled test")
    m = metrics["combined"]
    table([["Actual observation", "System flagged anomaly", "System did not flag anomaly"],
           ["Anomaly", f"{m['TP']} correct detections", f"{m['FN']} missed anomalies"],
           ["Normal", f"{m['FP']} false alarms", f"{m['TN']} correct normal results"]], [2.0, 2.25, 2.25])
    caption("Table 2 Performance measures")
    table([["Measure", "Result", "Meaning"],
           ["Precision", f"{m['precision']:.2%}", "How many alerts matched the introduced anomalies"],
           ["Recall", f"{m['recall']:.2%}", "How many introduced anomalies were detected"],
           ["F1 score", f"{m['f1']:.2%}", "A combined measure of precision and recall"],
           ["Accuracy", f"{m['accuracy']:.2%}", "How many observations were classified correctly overall"]], [1.3, .9, 4.3])
    heading("Discussion", 1)
    p("The system detected most of the introduced anomalies, but it also made mistakes. Recall of 75.00 percent means that "
      "it found three out of every four anomalies in this dataset. Precision of 63.16 percent shows that some alerts were false alarms.")
    p("Some valid rainfall differences between stations were incorrectly flagged. Some introduced errors, including smaller changes "
      "and partly reduced readings, were missed. These results show why a flagged observation should be reviewed before it is treated "
      "as a confirmed sensor problem.")
    p("The accuracy was 95.42 percent, but most observations in the dataset were normal. Accuracy should therefore be read together "
      "with precision and recall, rather than used alone to describe detection performance.")
    p("These findings apply only to this controlled test. The locations and rainfall readings were fictional, and the examples were created "
      "with knowledge of the system. The test does not establish accuracy on real PAGASA records. "
      "A separate evaluation using real neighboring-station data with independently checked labels is needed to make that claim.")
    destination = OUT / "Rainfall_Anomaly_Evaluation_IMRaD_Simple.docx"
    doc.save(destination)
    reopened = Document(destination)
    headings = [paragraph.text for paragraph in reopened.paragraphs if paragraph.style.name == "Heading 1"]
    assert headings == ["Introduction", "Methods", "Results", "Discussion"]
    assert len(reopened.tables) == 2
    text = "\n".join(paragraph.text for paragraph in reopened.paragraphs) + "\n" + "\n".join(cell.text for t in reopened.tables for row in t.rows for cell in row.cells)
    for expected in ("63.16%", "75.00%", "68.57%", "95.42%", "224", "217 correct normal results"):
        assert expected in text, expected
    print(f"Created {destination}; IMRaD heading order and two editable tables checked")


if __name__ == "__main__":
    main()
